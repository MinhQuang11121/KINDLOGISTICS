import pandas as pd
import numpy as np
from openpyxl import Workbook
from openpyxl.styles import Font, Alignment, PatternFill, Border, Side
from openpyxl.formatting.rule import ColorScaleRule

def create_scenario_data():
    file_path = "Kich_ban_Du_bao_KINDLOGISTICS.xlsx"
    
    # ==========================================
    # 1. TẠO DỮ LIỆU TỔNG QUAN (SHEET 1)
    # ==========================================
    summary_data = {
        "Thông số kịch bản": [
            "Tên kịch bản", 
            "Thời điểm giả lập", 
            "Điều kiện ngoại cảnh", 
            "Khu vực trọng điểm", 
            "Mục tiêu phân tích"
        ],
        "Chi tiết": [
            "Dự báo biến động nhu cầu khi mưa lớn cục bộ",
            "17:00 - 18:00 (Giờ cao điểm chiều)",
            "Lượng mưa 50mm, có dấu hiệu ngập tại các vùng trũng",
            "Quận Thanh Khê (Hàm Nghi, Thạc Gián), Quận Hải Châu",
            "Kiểm chứng độ nhạy của mô hình STGCN Đa đặc trưng với biến Mưa/Ngập"
        ]
    }
    df_summary = pd.DataFrame(summary_data)

    # ==========================================
    # 2. TẠO DỮ LIỆU SO SÁNH (SHEET 2)
    # ==========================================
    comparison_data = {
        "Khu vực (Nút giao)": [
            "Ngã tư Hàm Nghi - Hùng Vương", 
            "Ngã tư Thạc Gián", 
            "Nguyễn Văn Linh - Hoàng Diệu", 
            "Lê Duẩn - Ông Ích Khiêm", 
            "Điện Biên Phủ"
        ],
        "Lưu lượng thực tế (Bình thường)": [85, 42, 120, 95, 110],
        "Lưu lượng dự báo (Khi mưa lớn)": [132, 78, 145, 112, 128],
        "Tỷ lệ biến động (%)": [55.3, 85.7, 20.8, 17.9, 16.4],
        "Mức độ ngập giả lập (cm)": [15, 20, 5, 0, 10]
    }
    df_comp = pd.DataFrame(comparison_data)

    # ==========================================
    # 3. TẠO DỮ LIỆU CHUỖI THỜI GIAN (SHEET 3)
    # ==========================================
    time_series = {
        "Khung giờ": ["17:00", "17:15", "17:30", "17:45", "18:00"],
        "Lượng mưa (mm)": [5, 15, 45, 55, 30], # Lượng mưa tăng đột biến lúc 17:30
        "Nhu cầu thực tế (Lượt đơn)": [90, 105, 125, 140, 130],
        "Dự báo STGCN (Cơ bản)": [88, 102, 115, 120, 125],
        "Dự báo STGCN (Đa đặc trưng)": [92, 108, 128, 142, 135]
    }
    df_time = pd.DataFrame(time_series)

    # ==========================================
    # 4. XUẤT RA EXCEL VÀ TRANG TRÍ (FORMATTING)
    # ==========================================
    with pd.ExcelWriter(file_path, engine='openpyxl') as writer:
        df_summary.to_excel(writer, sheet_name='Tong_quan_Kich_ban', index=False)
        df_comp.to_excel(writer, sheet_name='So_sanh_Bien_dong', index=False)
        df_time.to_excel(writer, sheet_name='Du_bao_Theo_Thoi_gian', index=False)

        # Lấy file excel đang mở để format
        workbook = writer.book
        
        # Định dạng màu nền và viền cho Header
        header_fill = PatternFill(start_color="2C3E50", end_color="2C3E50", fill_type="solid")
        header_font = Font(color="FFFFFF", bold=True)
        thin_border = Border(left=Side(style='thin', color='BDC3C7'), 
                             right=Side(style='thin', color='BDC3C7'), 
                             top=Side(style='thin', color='BDC3C7'), 
                             bottom=Side(style='thin', color='BDC3C7'))
        
        for sheet_name in writer.sheets:
            ws = writer.sheets[sheet_name]
            
            # Tự động căn chỉnh độ rộng cột cho đẹp
            for col in ws.columns:
                max_length = 0
                column = col[0].column_letter
                for cell in col:
                    try:
                        if len(str(cell.value)) > max_length:
                            max_length = len(str(cell.value))
                    except:
                        pass
                ws.column_dimensions[column].width = max_length + 5

            # Tô màu Header
            for cell in ws[1]:
                cell.fill = header_fill
                cell.font = header_font
                cell.alignment = Alignment(horizontal='center', vertical='center')
                cell.border = thin_border
            
            # Kẻ bảng cho các dòng dữ liệu
            for row in ws.iter_rows(min_row=2):
                for cell in row:
                    cell.border = thin_border
                    cell.alignment = Alignment(horizontal='left', vertical='center')

        # Thêm rules tạo màu Heatmap tự động cho Sheet 2 (Cột D: % biến động, Cột E: Mức ngập)
        ws_comp = writer.sheets['So_sanh_Bien_dong']
        color_rule = ColorScaleRule(start_type='min', start_color='F9EBEA',
                                    end_type='max', end_color='E74C3C') # Đỏ dần khi số càng lớn
        ws_comp.conditional_formatting.add(f'D2:D{len(df_comp)+1}', color_rule) 
        ws_comp.conditional_formatting.add(f'E2:E{len(df_comp)+1}', color_rule) 

    print(f"✅ Đã tạo thành công file: {file_path}")

if __name__ == "__main__":
    create_scenario_data()